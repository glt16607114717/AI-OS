"""
Excel 读取工具

读取 Excel 文件内容，输出为 JSON 格式。

用法：python excel_read.py <Excel文件> [选项]
示例：python excel_read.py data.xlsx
      python excel_read.py data.xlsx --sheet Sheet1
      python excel_read.py data.xlsx --limit 10
      python excel_read.py data.xlsx --headers-only

选项：
  --sheet <名称>     指定工作表（默认第一个）
  --limit <数量>     限制读取行数（默认全部）
  --headers-only     只输出表头信息
"""
import json
import sys
import os

try:
    from openpyxl import load_workbook
except ImportError:
    print("[错误] 缺少 openpyxl 库，请安装: uv pip install openpyxl")
    sys.exit(1)


def main():
    if len(sys.argv) < 2:
        print("用法: python excel_read.py <Excel文件> [选项]")
        print("选项: --sheet <名称>     指定工作表（默认第一个）")
        print("      --limit <数量>     限制读取行数")
        print("      --headers-only     只输出表头")
        print("示例: python excel_read.py data.xlsx")
        print("      python excel_read.py data.xlsx --sheet Sheet1 --limit 10")
        sys.exit(1)

    args = sys.argv[1:]

    # 解析参数
    file_path = args[0]
    sheet_name = None
    limit = None
    headers_only = False

    i = 1
    while i < len(args):
        if args[i] == "--sheet" and i + 1 < len(args):
            sheet_name = args[i + 1]
            i += 2
        elif args[i] == "--limit" and i + 1 < len(args):
            limit = int(args[i + 1])
            i += 2
        elif args[i] == "--headers-only":
            headers_only = True
            i += 1
        else:
            i += 1

    if not os.path.isfile(file_path):
        print(f"[错误] 文件不存在: {file_path}")
        sys.exit(1)

    # 读取工作簿
    wb = load_workbook(file_path, read_only=True, data_only=True)

    # 获取工作表
    sheet_names = wb.sheetnames
    if sheet_name:
        if sheet_name not in sheet_names:
            print(f"[错误] 工作表不存在: {sheet_name}")
            print(f"可用工作表: {', '.join(sheet_names)}")
            wb.close()
            sys.exit(1)
        ws = wb[sheet_name]
    else:
        ws = wb[sheet_names[0]]
        sheet_name = sheet_names[0]

    # 读取所有行
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        result = {
            "success": True,
            "file": file_path,
            "sheet": sheet_name,
            "sheets": sheet_names,
            "headers": [],
            "data": [],
            "total_rows": 0,
        }
        print(json.dumps(result, ensure_ascii=False))
        return

    # 第一行为表头
    headers = [str(h) if h is not None else "" for h in rows[0]]
    data_rows = rows[1:]

    if headers_only:
        result = {
            "success": True,
            "file": file_path,
            "sheet": sheet_name,
            "sheets": sheet_names,
            "headers": headers,
            "total_rows": len(data_rows),
        }
        print(json.dumps(result, ensure_ascii=False))
        return

    # 限制行数
    if limit and limit < len(data_rows):
        data_rows = data_rows[:limit]

    # 转为字典列表
    data = []
    for row in data_rows:
        row_dict = {}
        for col_idx, header in enumerate(headers):
            value = row[col_idx] if col_idx < len(row) else None
            row_dict[header] = value
        data.append(row_dict)

    result = {
        "success": True,
        "file": os.path.basename(file_path),
        "sheet": sheet_name,
        "sheets": sheet_names,
        "headers": headers,
        "data": data,
        "total_rows": len(rows) - 1,
        "returned_rows": len(data),
    }
    print(json.dumps(result, ensure_ascii=False, default=str))


if __name__ == "__main__":
    main()
