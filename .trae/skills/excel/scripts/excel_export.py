"""
Excel 导出工具

将 JSON 数据导出为 Excel 文件。

用法：python excel_export.py <配置文件>
示例：python excel_export.py export_config.json

配置文件格式：
{
    "filename": "导出文件名（可选，默认"导出文件"）",
    "headers": ["列名1", "列名2", ...],
    "data": [
        {"列名1": "值1", "列名2": "值2"},
        ...
    ]
}
"""
import json
import sys
import os

try:
    from openpyxl import Workbook
except ImportError:
    print("[错误] 缺少 openpyxl 库，请安装: uv pip install openpyxl")
    sys.exit(1)


def main():
    if len(sys.argv) < 2:
        print("用法: python excel_export.py <配置文件>")
        print("示例: python excel_export.py export_config.json")
        print("")
        print("配置文件格式：")
        print(json.dumps({
            "filename": "导出文件名（可选）",
            "headers": ["列名1", "列名2"],
            "data": [
                {"列名1": "值1", "列名2": "值2"},
            ]
        }, ensure_ascii=False, indent=4))
        sys.exit(1)

    config_file = sys.argv[1]

    with open(config_file, "r", encoding="utf-8") as f:
        config = json.load(f)

    headers = config.get("headers", [])
    data = config.get("data", [])
    filename = config.get("filename", "导出文件")

    if not headers:
        print("[错误] headers 不能为空")
        sys.exit(1)

    if not data:
        print("[错误] data 不能为空")
        sys.exit(1)

    # 确保 filename 以 .xlsx 结尾
    if not filename.endswith(".xlsx"):
        filename += ".xlsx"

    # 输出目录为配置文件所在目录
    output_dir = os.path.dirname(os.path.abspath(config_file))
    output_path = os.path.join(output_dir, filename)

    # 创建工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"

    # 写入表头
    for col_idx, header in enumerate(headers, 1):
        from openpyxl.styles import Font
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)

    # 写入数据
    for row_idx, row_data in enumerate(data, 2):
        for col_idx, header in enumerate(headers, 1):
            value = row_data.get(header, "")
            ws.cell(row=row_idx, column=col_idx, value=value)

    # 自动调整列宽
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                cell_length = len(str(cell.value or ""))
                if cell_length > max_length:
                    max_length = cell_length
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 4, 50)

    wb.save(output_path)

    result = {
        "success": True,
        "filename": filename,
        "path": output_path,
        "rows": len(data),
        "columns": len(headers),
    }
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
